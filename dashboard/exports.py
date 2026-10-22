"""Daily order exports for the staff dashboard."""
from datetime import datetime, time, timedelta
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone

from orders.models import Order


def daily_orders(day):
    start = timezone.make_aware(datetime.combine(day, time.min))
    end = timezone.make_aware(datetime.combine(day + timedelta(days=1), time.min))
    return (
        Order.objects.filter(created_at__gte=start, created_at__lt=end)
        .prefetch_related("items")
        .order_by("created_at")
    )


def _rows(orders):
    rows = [["Order", "Date/time", "Customer", "Email", "Phone", "Items", "Payment", "Payment status", "Fulfillment", "Subtotal", "Discount", "Shipping", "Total"]]
    for order in orders:
        item_summary = "; ".join(
            f"{item.product_name} ({item.size}/{item.color}) x{item.quantity}"
            for item in order.items.all()
        )
        rows.append([
            order.order_number,
            timezone.localtime(order.created_at).strftime("%Y-%m-%d %H:%M"),
            order.full_name,
            order.email,
            order.phone,
            item_summary,
            order.get_payment_method_display(),
            order.get_payment_status_display(),
            order.get_status_display(),
            str(order.subtotal),
            str(order.discount_amount),
            str(order.shipping_amount),
            str(order.total_amount),
        ])
    return rows


def export_daily_xlsx(day):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Orders"
    for row in _rows(daily_orders(day)):
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="25382E")
    for column in sheet.columns:
        width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 55)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = width
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="orders-{day.isoformat()}.xlsx"'
    return response


def export_daily_pdf(day):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

    rows = _rows(daily_orders(day))
    # Keep rendered table compact and printable while retaining each order's line items.
    compact = [[row[0], row[1], row[2], row[3], row[5], row[6], row[7], row[8], row[12]] for row in rows]
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(letter), title=f"Orders {day.isoformat()}")
    styles = getSampleStyleSheet()
    data = [[Paragraph(str(value), styles["BodyText"]) for value in row] for row in compact]
    table = Table(data, repeatRows=1, colWidths=[70, 75, 75, 105, 170, 65, 65, 65, 55])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#25382e")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d9ded8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f6f3")]),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
    ]))
    doc.build([Paragraph(f"Daily orders — {day.isoformat()}", styles["Title"]), table])
    response = HttpResponse(output.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="orders-{day.isoformat()}.pdf"'
    return response
